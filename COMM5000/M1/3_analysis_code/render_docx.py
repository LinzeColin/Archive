import re, content as C
from docx import Document
from docx.shared import Pt, Mm, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

SERIF, SANS = 'Georgia', 'Arial'   # universally available on Mac/Windows Word
INK = RGBColor(0x1f, 0x1f, 0x1d); BLUE = RGBColor(0x2a, 0x78, 0xd6); ORANGE = RGBColor(0xc4, 0x50, 0x1f); GREY = RGBColor(0x52, 0x51, 0x4e)
def mn(t): return re.sub(r'(?<![\w\-])-(?=[\d.])', '−', str(t))

doc = Document()
sec = doc.sections[0]
sec.page_height, sec.page_width = Mm(297), Mm(210)
sec.left_margin = sec.right_margin = Mm(16); sec.top_margin = Mm(15); sec.bottom_margin = Mm(16)
st = doc.styles['Normal']; st.font.name = SERIF; st.font.size = Pt(9.5); st.element.rPr.rFonts.set(qn('w:eastAsia'), SERIF)
st.paragraph_format.space_after = Pt(4); st.paragraph_format.line_spacing = 1.18

def shade(cell, hexcol):
    tcPr = cell._tc.get_or_add_tcPr(); s = OxmlElement('w:shd'); s.set(qn('w:val'), 'clear'); s.set(qn('w:color'), 'auto'); s.set(qn('w:fill'), hexcol); tcPr.append(s)

def borders(table):
    tbl = table._tbl; tblPr = tbl.tblPr
    b = OxmlElement('w:tblBorders')
    for edge, sz, col in [('top', 8, '1f1f1d'), ('bottom', 8, '1f1f1d'), ('insideH', 2, 'd9d8d3'), ('left', 0, 'ffffff'), ('right', 0, 'ffffff'), ('insideV', 0, 'ffffff')]:
        e = OxmlElement(f'w:{edge}'); e.set(qn('w:val'), 'single' if sz else 'nil'); e.set(qn('w:sz'), str(sz)); e.set(qn('w:color'), col); b.append(e)
    tblPr.append(b)

def rich(par, text, size=None, font=None, color=None):
    parts = re.split(r'(<b>.*?</b>|<i>.*?</i>)', text)
    for part in parts:
        if not part: continue
        bold = part.startswith('<b>'); ital = part.startswith('<i>')
        t = re.sub(r'</?[bi]>', '', part)
        r = par.add_run(mn(t)); r.bold = bold; r.italic = ital
        if size: r.font.size = Pt(size)
        if font: r.font.name = font
        if color: r.font.color.rgb = color

def small(par_text, size=7.2, color=GREY, bold_lead=None):
    p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(6)
    if bold_lead:
        r = p.add_run(bold_lead + ' '); r.bold = True; r.font.size = Pt(size); r.font.name = SANS
    r = p.add_run(mn(par_text)); r.font.size = Pt(size); r.font.name = SANS; r.font.color.rgb = color
    return p

def table(head, rows, widths=None, size=7.0, numeric_from=None, seg_col=None):
    t = doc.add_table(rows=1, cols=len(head)); t.alignment = WD_TABLE_ALIGNMENT.CENTER; t.autofit = False; borders(t)
    for j, h in enumerate(head):
        c = t.rows[0].cells[j]; c.text = ''; shade(c, 'f4f3ef')
        p = c.paragraphs[0]; r = p.add_run(h); r.bold = True; r.font.size = Pt(size); r.font.name = SANS
        if numeric_from is not None and j >= numeric_from: p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for row in rows:
        cells = t.add_row().cells
        for j, v in enumerate(row):
            c = cells[j]; c.text = ''; p = c.paragraphs[0]; p.paragraph_format.space_after = Pt(0)
            r = p.add_run(mn(v)); r.font.size = Pt(size); r.font.name = SANS
            if numeric_from is not None and j >= numeric_from: p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            if seg_col is not None and j == seg_col:
                if v == 'Luxury': r.font.color.rgb = ORANGE; r.bold = True
                elif v.startswith('Non'): r.font.color.rgb = BLUE
            if j == 0 and seg_col is not None and v: r.bold = True
    if widths:
        for row in t.rows:
            for j, w in enumerate(widths): row.cells[j].width = Mm(w)
    return t

def caption(text):
    p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(6); p.paragraph_format.space_after = Pt(2); p.paragraph_format.keep_with_next = True
    r = p.add_run(text); r.bold = True; r.font.size = Pt(7.8); r.font.name = SANS

wc = sum(len(re.sub(r'<[^>]+>', '', b[1]).split()) for b in C.BODY if b[0] in ('p', 'h'))
p = doc.add_paragraph(); r = p.add_run(C.META.upper()); r.font.size = Pt(7); r.font.name = SANS; r.font.color.rgb = GREY
r = p.add_run(f"     Body text: {wc} words (excl. tables, figures, captions, references)"); r.font.size = Pt(7); r.font.name = SANS; r.font.color.rgb = GREY
p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(0); r = p.add_run(C.TITLE); r.bold = True; r.font.size = Pt(22); r.font.name = SANS; r.font.color.rgb = INK
p = doc.add_paragraph(); r = p.add_run(C.SUBTITLE); r.font.size = Pt(11); r.font.name = SANS; r.font.color.rgb = GREY
small('online used-car marketplace, India   ·   Data: COMM5000-Used_Car_Price_Prediction.xlsm · 1,005,000 listings × 20 fields · full-population analysis', bold_lead='Client')

for b in C.BODY:
    if b[0] == 'h':
        n, t = b[1].split('  ', 1)
        p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(9); p.paragraph_format.space_after = Pt(3); p.paragraph_format.keep_with_next = True
        r = p.add_run(n + '  '); r.bold = True; r.font.size = Pt(11.5); r.font.name = SANS; r.font.color.rgb = BLUE
        r = p.add_run(t); r.bold = True; r.font.size = Pt(11.5); r.font.name = SANS; r.font.color.rgb = INK
    elif b[0] == 'p':
        p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY; rich(p, b[1])
    elif b[0] == 'kpi':
        t = doc.add_table(rows=1, cols=4); t.alignment = WD_TABLE_ALIGNMENT.CENTER
        for j, (v, l, s) in enumerate(C.KPIS):
            c = t.rows[0].cells[j]; shade(c, 'f6f5f2'); c.text = ''
            p = c.paragraphs[0]; r = p.add_run(mn(v)); r.bold = True; r.font.size = Pt(13); r.font.name = SANS; r.font.color.rgb = ORANGE if j == 3 else BLUE
            p = c.add_paragraph(); r = p.add_run(l); r.font.size = Pt(7); r.font.name = SANS
            p = c.add_paragraph(); r = p.add_run(mn(s)); r.font.size = Pt(6.5); r.font.name = SANS; r.font.color.rgb = GREY
        doc.add_paragraph().paragraph_format.space_after = Pt(0)
    elif b[0] == 'fig':
        path, cap = C.FIGS[b[1]]
        p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.keep_with_next = True
        p.add_run().add_picture(path, width=Mm(176))
        lab, rest = cap.split('. ', 1); small(rest, bold_lead=lab + '.')
    elif b[0] == 'table':
        if b[1] == 'audit':
            caption('Table 1. Data-quality audit and treatment log'); table(C.AUDIT_HEAD, C.AUDIT, [38, 32, 38, 28, 42], 6.8)
        elif b[1] == 'desc':
            caption('Table 2. Numerical summaries: entire sample, luxury (Audi, BMW, Mercedes) and non-luxury segments')
            table(C.DESC_HEAD, C.desc_rows(), [36, 14, 18, 18, 18, 16, 18, 16, 20], 6.8, numeric_from=2, seg_col=1); small(C.DESC_NOTE, 6.5)
        elif b[1] == 'plan':
            caption('Table 3. Development plan for Milestone 2 and the final report'); table(C.PLAN_HEAD, C.PLAN, [34, 14, 72, 58], 6.8)

p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(9)
r = p.add_run('References'); r.bold = True; r.font.size = Pt(11.5); r.font.name = SANS
for ref in C.REFS:
    p = doc.add_paragraph(); p.paragraph_format.left_indent = Cm(0.5); p.paragraph_format.first_line_indent = Cm(-0.5); rich(p, ref, size=8.6)
p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(9)
r = p.add_run('Appendix: robustness of key results'); r.bold = True; r.font.size = Pt(11.5); r.font.name = SANS
caption('Table A1a. Sensitivity to cleaning choices'); table(C.SENS_HEAD, C.sens_rows(), [46, 22, 28, 24, 26, 26], 6.8, numeric_from=1)
caption('Table A1b. Usage correlation under alternative km caps'); table(["Kms_Driven cap", "n", "r (raw)", "r (log–log)"], C.cap_rows(), [50, 30, 30, 30], 6.8, numeric_from=1)
small('All statistics computed on the full saved dataset (no sampling) except the random-sample panels (Figure 2B; Figure 3A, 3C). Confidence intervals for medians use order statistics; for the segment median ratio, 200 bootstrap resamples. The companion Excel workbook lists every aggregate, the cleaning maps and the source row numbers of plotted samples.', 6.8)

# footer page numbers
fp = sec.footer.paragraphs[0]; fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
r = fp.add_run('COMM5000 Milestone 1 · page '); r.font.size = Pt(7); r.font.name = SANS; r.font.color.rgb = GREY
for kind, txt in [('begin', None), (None, 'PAGE'), ('end', None)]:
    run = fp.add_run(); run.font.size = Pt(7)
    if kind:
        e = OxmlElement('w:fldChar'); e.set(qn('w:fldCharType'), kind); run._r.append(e)
    else:
        e = OxmlElement('w:instrText'); e.set(qn('xml:space'), 'preserve'); e.text = txt; run._r.append(e)
# anonymous metadata
cp = doc.core_properties; cp.author = ''; cp.last_modified_by = ''; cp.title = C.TITLE; cp.comments = ''
doc.save('report.docx'); print('docx ok, words', wc)
