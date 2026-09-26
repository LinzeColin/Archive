import pandas as pd, numpy as np, json, math
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import content as C

df = pd.read_pickle('clean.pkl'); R = json.load(open('results.json')); PERM = json.load(open('perm.json'))
wb = Workbook()
H = Font(bold=True, color='FFFFFF', name='Arial', size=9); HF = PatternFill('solid', fgColor='1F1F1D')
B = Font(name='Arial', size=9); BB = Font(name='Arial', size=9, bold=True); T = Font(name='Arial', size=13, bold=True)
NOTE = Font(name='Arial', size=8, italic=True, color='52514E'); thin = Side(style='thin', color='D9D8D3')

def sheet(name, title, head, rows, widths=None, note=None, fmt=None):
    ws = wb.create_sheet(name)
    ws['A1'] = title; ws['A1'].font = T
    r0 = 3
    if note:
        ws['A2'] = note; ws['A2'].font = NOTE
    for j, h in enumerate(head, 1):
        c = ws.cell(r0, j, h); c.font = H; c.fill = HF; c.alignment = Alignment(wrap_text=True, vertical='center')
    for i, row in enumerate(rows, r0 + 1):
        for j, v in enumerate(row, 1):
            c = ws.cell(i, j, v); c.font = B; c.border = Border(bottom=thin)
            if fmt and j in fmt and isinstance(v, (int, float)): c.number_format = fmt[j]
    for j, w in enumerate(widths or [], 1): ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = ws.cell(r0 + 1, 1)
    return ws

# 0 README
ws = wb.active; ws.title = 'README'
lines = [
 ('COMM5000 M1 — companion analysis workbook', T),
 ('Purpose: every number in the report can be traced here. Source: COMM5000-Used_Car_Price_Prediction.xlsm, sheet DATA (A1:T1005001), Dataset!B6 stamp 2026-09-26 08:29:38.', B),
 ('The source workbook was read only; no macro was run and B6 was not changed. All statistics use the full 1,005,000 rows unless a sheet says otherwise.', B),
 ('', B),
 ('Sheets', BB),
 ('Cleaning_Log — every data issue, detection method, count, treatment and rationale (Report Table 1).', B),
 ('Label_Map — each raw Brand / Fuel_Type label, its cleaned value and frequency.', B),
 ('Descriptive_Stats — Report Table 2 with extra columns (Q1, Q3, skewness, CV).', B),
 ('Correlation_Inputs — n, Σx, Σy, Σx², Σy², Σxy per group; Pearson r is recomputed by live Excel formulas and matches the report.', B),
 ('Age_Profile / Km_Deciles / Brand_Medians — full-data group summaries behind Figures 3 and 4.', B),
 ('Segment_Test — luxury vs non-luxury tests (Welch t, Mann–Whitney, effect size, bootstrap CI).', B),
 ('Signal_Scan — variance share of log price explained by each of 19 fields (Figure 4B).', B),
 ('Duplicate_Permutation — 20 column-shuffle runs giving the chance baseline for 15-field repeats.', B),
 ('Sensitivity — Appendix Table A1 (cleaning choices, km caps).', B),
 ('Sampling_r — 1,000 random 5,000-row correlations (Figure 2B, summary).', B),
 ('Plot_Sample_Age / Plot_Sample_Km — the exact 5,000 rows plotted in Figure 3A and 3C, with source Excel row numbers.', B),
 ('Excel_Recipes — the Excel formulas that reproduce the cleaning and key statistics directly on the DATA sheet.', B),
]
for i, (t, f) in enumerate(lines, 1):
    ws.cell(i, 1, t).font = f
ws.column_dimensions['A'].width = 140

# 1 Cleaning log
sheet('Cleaning_Log', 'Data-quality audit and treatment log (Report Table 1)', C.AUDIT_HEAD, C.AUDIT, [42, 34, 40, 30, 46])
tl = wb['Cleaning_Log']; r = 4 + len(C.AUDIT) + 1
for k, v in [('Technical-value rule thresholds (mean ± 3 SD, zeros excluded)', '')] + [(f"{c}: zeros {R['tech'][c]['zero']}, beyond ±3 SD {R['tech'][c]['sd3']}, bounds {R['tech'][c]['lo']:.2f} – {R['tech'][c]['hi']:.2f}, excess kurtosis {R['tech'][c]['kurt']:.3f}", '') for c in R['tech']]:
    tl.cell(r, 1, k).font = B; r += 1

# 2 Label map
rows = []
for col, clean in [('Brand', 'BrandC'), ('Fuel_Type', 'FuelC')]:
    vc = df.groupby([col, clean]).size().reset_index(name='n')
    for _, x in vc.iterrows():
        rows.append([col, repr(x[col]), x[clean], int(x['n']), 'changed' if x[col] != x[clean] else 'unchanged', 'Luxury' if (col == 'Brand' and x[clean] in ('Audi', 'BMW', 'Mercedes')) else ''])
sheet('Label_Map', 'Raw label → cleaned label (quotes show hidden spaces)', ['Field', 'Raw label', 'Cleaned', 'Rows', 'Status', 'Segment'], rows, [12, 18, 14, 10, 12, 10], fmt={4: '#,##0'})

# 3 Descriptive stats
rows = []
for v, lab, dp in C.DESC_VARS:
    for g in ['All', 'Luxury', 'Non-luxury']:
        d = R['desc'][g][v]
        rows.append([lab, g, d['n'], d['mean'], d['median'], d['mode'], d['sd'], d['min'], d['max'], d['q1'], d['q3'], d['skew'], d['cv']])
sheet('Descriptive_Stats', 'Numerical summaries (Report Table 2, extended)', ['Variable', 'Group', 'n', 'Mean', 'Median', 'Mode*', 'SD (n−1)', 'Min', 'Max', 'Q1', 'Q3', 'Skewness', 'CV'], rows,
      [24, 12, 11, 14, 14, 12, 14, 12, 14, 14, 14, 10, 8], note=C.DESC_NOTE, fmt={3: '#,##0', **{j: '#,##0.00' for j in range(4, 12)}, 12: '0.000', 13: '0.000'})

# 4 Correlation inputs with live formulas
rows = []
for g, sub in [('All', df), ('Luxury', df[df.Segment == 'Luxury']), ('Non-luxury', df[df.Segment == 'Non-luxury'])]:
    for xname, xcol in [('Registration_Age', 'Registration_Age'), ('Kms_Driven ≤ 1m', 'Kms_c'), ('Kms_Driven (all)', 'Kms_Driven')]:
        s = sub[[xcol, 'Car_Price']].dropna(); x = s[xcol].astype(float).values; y = s.Car_Price.values
        rows.append([g, xname, len(s), math.fsum(x), math.fsum(y), math.fsum(x * x), math.fsum(y * y), math.fsum(x * y), None, R['rel'][g][{'Registration_Age': 'age', 'Kms_c': 'km', 'Kms_Driven': 'km_all'}[xcol]]['pearson']])
ws = sheet('Correlation_Inputs', 'Pearson r from sufficient statistics (column I is a live Excel formula)', ['Group', 'x variable', 'n', 'Σx', 'Σy (price)', 'Σx²', 'Σy²', 'Σxy', 'r (Excel formula)', 'r (Python, full data)'], rows,
           [12, 18, 11, 20, 22, 22, 26, 24, 16, 18], note='r = (nΣxy − ΣxΣy) / √[(nΣx² − (Σx)²)(nΣy² − (Σy)²)]', fmt={3: '#,##0', 4: '0.000E+00', 5: '0.000E+00', 6: '0.000E+00', 7: '0.000E+00', 8: '0.000E+00', 10: '0.000000'})
for i in range(4, 4 + len(rows)):
    ws.cell(i, 9, f'=(C{i}*H{i}-D{i}*E{i})/SQRT((C{i}*F{i}-D{i}^2)*(C{i}*G{i}-E{i}^2))').number_format = '0.000000'

# 5 Age profile
def med_ci(x):
    x = np.sort(np.asarray(x)); n = len(x); k = 1.96 * math.sqrt(n) / 2
    return np.median(x), x[max(int(n / 2 - k), 0)], x[min(int(n / 2 + k), n - 1)]
rows = []
for a, g in df.groupby('Registration_Age'):
    row = [int(a)]
    for seg in ['All', 'Luxury', 'Non-luxury']:
        v = g.Car_Price.values if seg == 'All' else g.Car_Price[g.Segment == seg].values
        m, lo, hi = med_ci(v); row += [len(v), float(np.mean(v)), float(m), float(lo), float(hi)]
    rows.append(row)
head = ['Age']
for seg in ['All', 'Lux', 'Non-lux']: head += [f'{seg} n', f'{seg} mean', f'{seg} median', f'{seg} median CI lo', f'{seg} median CI hi']
sheet('Age_Profile', 'Car_Price by Registration_Age, full data (Figure 3B)', head, rows, [6] + [11] * 15, fmt={j: '#,##0' for j in range(2, 17)})

# 6 Km deciles
kv = df.dropna(subset=['Kms_c']).copy(); kv['dec'] = pd.qcut(kv.Kms_c, 10, labels=False)
rows = []
for d_, g in kv.groupby('dec'):
    row = [f'D{d_+1}', float(g.Kms_c.min()), float(g.Kms_c.max())]
    for seg in ['All', 'Luxury', 'Non-luxury']:
        v = g.Car_Price.values if seg == 'All' else g.Car_Price[g.Segment == seg].values
        m, lo, hi = med_ci(v); row += [len(v), float(m), float(lo), float(hi)]
    rows.append(row)
head = ['Decile', 'km from', 'km to']
for seg in ['All', 'Lux', 'Non-lux']: head += [f'{seg} n', f'{seg} median', f'{seg} CI lo', f'{seg} CI hi']
sheet('Km_Deciles', 'Car_Price by Kms_Driven decile (≤ 1,000,000 km), full data (Figure 3D)', head, rows, [8, 12, 12] + [11] * 12, fmt={j: '#,##0' for j in range(2, 16)})

# 7 Brand medians
rows = []
for b_, g in df.groupby('BrandC'):
    m, lo, hi = med_ci(g.Car_Price.values); rows.append([b_, 'Luxury' if b_ in ('Audi', 'BMW', 'Mercedes') else 'Non-luxury', len(g), float(g.Car_Price.mean()), float(m), float(lo), float(hi)])
sheet('Brand_Medians', 'Car_Price by cleaned brand (Figure 4A)', ['Brand', 'Segment', 'n', 'Mean', 'Median', '95% CI lo', '95% CI hi'], sorted(rows, key=lambda r: r[4]), [14, 12, 10, 14, 14, 14, 14], fmt={3: '#,##0', 4: '#,##0', 5: '#,##0', 6: '#,##0', 7: '#,##0'})

# 8 Segment test
S = R['seg']
rows = [['Mean difference (Lux − Non-lux), INR', S['mean_diff']], ['95% CI (Welch)', f"{S['mean_diff_ci'][0]:,.0f} to {S['mean_diff_ci'][1]:,.0f}"], ['Mean difference, %', S['mean_diff_pct']],
        ['Median difference, %', S['median_diff_pct']], ['Median ratio 95% bootstrap CI, %', f"{S['median_ratio_ci_pct'][0]:+.3f} to {S['median_ratio_ci_pct'][1]:+.3f}"],
        ['Welch t (raw price)', S['welch_t']], ['Welch p (raw price)', S['welch_p']], ['Welch p (log price)', S['log_welch_p']], ['Mann–Whitney p', S['mw_p']], ['Cohen d (log price)', S['cohen_d_log']],
        ['Luxury rows if brand labels NOT trimmed', R['lux_n_if_uncleaned']], ['Luxury rows after trimming', R['n_lux']]]
sheet('Segment_Test', 'Luxury vs non-luxury price comparison (Section 5)', ['Statistic', 'Value'], rows, [44, 26], fmt={2: '#,##0.0000'})

# 9 Signal scan
lab = {'Kms_c': 'Kms_Driven (log–log r²)', 'Horsepower_c': 'Horsepower (log–log r²)', 'Engine_CC_c': 'Engine_CC (log–log r²)', 'Mileage_kmpl_c': 'Mileage_kmpl (log–log r²)', 'BrandC': 'Brand (η²)', 'FuelC': 'Fuel_Type (η²)', 'Segment': 'Segment (η²)'}
rows = [[lab.get(k, k + ' (η²)'), v, v * 100] for k, v in sorted(R['screen_r2'].items(), key=lambda kv: -kv[1])]
sheet('Signal_Scan', 'Share of log-price variance explained by each field (Figure 4B)', ['Field (measure)', 'Share (0–1)', 'Share (%)'], rows, [30, 14, 12], fmt={2: '0.000000000', 3: '0.000000'})
wsg = wb['Signal_Scan']; n0 = 4 + len(rows) + 1
for k, v in [('Preview log-linear model (log km + log hp + log cc): R²', R['prev_model']['r2']), ('Preview model MdAPE, %', R['prev_model']['mdape']), ('MdAPE of quoting overall median, %', R['mdape_global_median']),
             ('MdAPE of segment medians, %', R['mdape_segment_median']), ('MdAPE of age medians, %', R['mdape_age_median']), ('MdAPE of 20 km-bin medians, %', R['mdape_km20_median']), ('Correlation Engine_CC vs Horsepower', R['corr_engine_hp'])]:
    wsg.cell(n0, 1, k).font = B; wsg.cell(n0, 2, v).font = B; wsg.cell(n0, 2).number_format = '0.0000'; n0 += 1

# 10 Duplicate permutation
rows = [[i + 1, v] for i, v in enumerate(PERM['nulls'])]
ws = sheet('Duplicate_Permutation', '15-field repeats: observed vs chance (each column shuffled independently, seed 1)', ['Run', 'Repeat rows after shuffling'], rows, [8, 26], fmt={2: '#,##0'})
n0 = 4 + len(rows) + 1
for k, f in [('Observed repeats', PERM['observed']), ('Mean of shuffles', f'=AVERAGE(B4:B{3+len(rows)})'), ('SD of shuffles', f'=STDEV.S(B4:B{3+len(rows)})'), ('Excess ≈ genuine duplicates', f'=B{n0}-B{n0+1}')]:
    ws.cell(n0, 1, k).font = BB; c = ws.cell(n0, 2, f); c.font = BB; c.number_format = '#,##0'; n0 += 1

# 11 Sensitivity
ws = sheet('Sensitivity', 'Appendix Table A1a/A1b', C.SENS_HEAD, C.sens_rows(), [30, 12, 18, 14, 16, 16])
n0 = 4 + 4 + 2
for j, h in enumerate(['Kms_Driven cap', 'n', 'r (raw)', 'r (log–log)'], 1):
    c = ws.cell(n0, j, h); c.font = H; c.fill = HF
for i, row in enumerate(C.cap_rows(), n0 + 1):
    for j, v in enumerate(row, 1): ws.cell(i, j, v).font = B

# 12 Sampling r
kvv = kv[['Kms_c', 'Car_Price']].values; av = df[['Registration_Age', 'Car_Price']].values; rng = np.random.default_rng(7); rows = []
for i in range(1000):
    ia = rng.integers(0, len(av), 5000); ik = rng.integers(0, len(kvv), 5000)
    rows.append([i + 1, float(np.corrcoef(av[ia, 0], av[ia, 1])[0, 1]), float(np.corrcoef(kvv[ik, 0], kvv[ik, 1])[0, 1])])
sheet('Sampling_r', '1,000 random 5,000-row samples (Figure 2B; same generator as the figure)', ['Draw', 'r price–age', 'r price–km'], rows, [8, 14, 14], fmt={2: '0.0000', 3: '0.0000'})

# 13-14 plot samples
for nm, f, cols in [('Plot_Sample_Age', 'plot_sample_age.csv', ['ExcelRow', 'Segment', 'Registration_Age', 'Car_Price']), ('Plot_Sample_Km', 'plot_sample_km.csv', ['ExcelRow', 'Segment', 'Kms_c', 'Car_Price'])]:
    s = pd.read_csv(f)
    sheet(nm, f'Exact rows plotted in Figure 3 ({"A" if "Age" in nm else "C"}); ExcelRow = row number on sheet DATA', ['DATA row', 'Segment', cols[2].replace('Kms_c', 'Kms_Driven'), 'Car_Price'], s[cols].values.tolist(), [10, 12, 16, 16], fmt={3: '#,##0.00', 4: '#,##0.00'})

# 15 Excel recipes
rec = [
 ['Clean brand', 'U2 (helper)', '=TRIM(A2)', 'Removes the padded spaces (10,040 rows).'],
 ['Clean fuel', 'V2 (helper)', '=LET(f,LOWER(TRIM(G2)),IF(f="hybridd","Hybrid",IF(f="electrik","Electric",IF(f="cng","CNG",PROPER(f)))))', 'Maps the 6 variants onto 5 fuel types + Unknown.'],
 ['Segment', 'W2 (helper)', '=IF(OR(U2="Audi",U2="BMW",U2="Mercedes"),"Luxury","Non-luxury")', 'Luxury = Audi, BMW, Mercedes.'],
 ['Usage flag', 'X2 (helper)', '=IF(L2>1000000,"exclude from km analysis","ok")', '6,134 rows flagged.'],
 ['Median price, luxury', 'any cell', '=MEDIAN(FILTER(DATA!T2:T1005001,DATA!W2:W1005001="Luxury"))', 'Returns 672,614.'],
 ['Mean price, all', 'any cell', '=AVERAGE(DATA!T2:T1005001)', 'Returns 1,015,107.'],
 ['Correlation price–age', 'any cell', '=CORREL(DATA!S2:S1005001,DATA!T2:T1005001)', 'Returns 0.0000672.'],
 ['Correlation price–km (≤1m)', 'any cell', '=CORREL(FILTER(DATA!L2:L1005001,DATA!L2:L1005001<=1000000),FILTER(DATA!T2:T1005001,DATA!L2:L1005001<=1000000))', 'Returns −0.0510.'],
 ['Random sample for scatter', 'helper', '=RAND() then sort and take first 5,000 rows', 'Or reuse the exact rows in Plot_Sample_* sheets.'],
 ['Blank check', 'any cell', '=COUNTBLANK(DATA!A2:T1005001)', 'Returns 0.'],
]
sheet('Excel_Recipes', 'Excel formulas that reproduce the cleaning and headline statistics on sheet DATA', ['Step', 'Where', 'Formula', 'Result / note'], rec, [26, 14, 90, 40])
wb.properties.creator = ''; wb.properties.lastModifiedBy = ''
wb.save('COMM5000_M1_Analysis_Workbook.xlsx'); print('xlsx ok')
