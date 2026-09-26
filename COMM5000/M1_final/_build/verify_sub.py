import openpyxl, pandas as pd, numpy as np, json, sys, math
from scipy import stats
path=sys.argv[1]; N=int(sys.argv[2])
wbv=openpyxl.load_workbook(path,data_only=True)
errs={}
for ws in wbv.worksheets:
    if ws.title in ('DATA','Dup_Sorted') and N>50000: continue
    for row in ws.iter_rows():
        for c in row:
            if isinstance(c.value,str) and c.value.startswith('#') and c.value[1:4].isupper():
                errs.setdefault(ws.title,[]).append((c.coordinate,c.value))
print('ERRORS:',{k:(len(v),v[:5]) for k,v in errs.items()})
raw=pd.read_pickle('raw.pkl').iloc[:N].copy()
fm={'petrol':'Petrol','diesel':'Diesel','cng':'CNG','electric':'Electric','electrik':'Electric','hybrid':'Hybrid','hybridd':'Hybrid','unknown':'Unknown'}
raw['BrandC']=raw.Brand.str.strip(); raw['FuelC']=raw.Fuel_Type.str.strip().str.lower().map(fm)
raw['Seg']=np.where(raw.BrandC.isin(['Audi','BMW','Mercedes']),'Luxury','Non-luxury'); raw['KMOK']=raw.Kms_Driven<=1e6
def ok(s):
    p=s[s>0]; return (s>0)&((s-p.mean()).abs()<=3*p.std(ddof=1))
raw['MilOK'],raw['EngOK'],raw['HPOK']=ok(raw.Mileage_kmpl),ok(raw.Engine_CC),ok(raw.Horsepower)
S=wbv['Summary']; M=json.load(open(path.split('/')[-1].replace('lo/','')+'.map.json')) if False else None
checks=[]
def chk(name,got,exp,tol=1e-6):
    good = (got is not None) and (abs(got-exp) <= tol*max(1,abs(exp)))
    checks.append((name,got,exp,good))
V={'Price':('Car_Price',None),'Age':('Registration_Age',None),'KM':('Kms_Driven','KMOK'),'Eng':('Engine_CC','EngOK'),'HP':('Horsepower','HPOK'),'Mil':('Mileage_kmpl','MilOK'),'Acc':('Accidents',None)}
r=5
for v in ['Price','Age','KM','Eng','HP','Mil','Acc']:
    col,flag=V[v]
    for g in ['All','Luxury','Non-luxury']:
        sub=raw if g=='All' else raw[raw.Seg==g]
        if flag: sub=sub[sub[flag]]
        x=sub[col]
        vals=[S.cell(r,k).value for k in range(3,13)]
        exp=[len(x),x.mean(),x.median(),None,x.std(ddof=1),x.min(),x.max(),x.quantile(.25),x.quantile(.75),stats.skew(x,bias=False)]
        for nm,gv,ev in zip(['n','mean','median','mode','sd','min','max','q1','q3','skew'],vals,exp):
            if ev is None: continue
            chk(f'Summary {v}/{g} {nm}',gv,float(ev))
        r+=1
C=wbv['Correlations']
cv=[('Registration_Age',None),('Year',None),('Kms_Driven','KMOK'),('Mileage_kmpl','MilOK'),('Engine_CC','EngOK'),('Horsepower','HPOK'),('Accidents',None),('Insurance_Valid',None),('Service_History',None),('Tax_Paid',None),('Number_of_Doors',None),('Seats',None)]
for i,(col,flag) in enumerate(cv):
    for j,g in enumerate(['All','Luxury','Non-luxury']):
        sub=raw if g=='All' else raw[raw.Seg==g]
        if flag: sub=sub[sub[flag]]
        chk(f'Corr {col}/{g}',C.cell(4+i,3+j).value,float(np.corrcoef(sub[col],sub.Car_Price)[0,1]))
k=raw[raw.KMOK]; lx,ly=np.log(k.Kms_Driven),np.log(k.Car_Price); b1,b0=np.polyfit(lx,ly,1)
chk('loglog r',C['B19'].value,np.corrcoef(lx,ly)[0,1]); chk('slope',C['B20'].value,b1); chk('intercept',C['B21'].value,b0)
chk('mdape model',C['B24'].value,np.median(np.abs(k.Car_Price-np.exp(b0+b1*lx))/k.Car_Price))
chk('mdape median',C['B25'].value,np.median(np.abs(raw.Car_Price-raw.Car_Price.median())/raw.Car_Price))
N_=wbv['Notes']
chk('brand relabel',N_['C14'].value,(raw.Brand!=raw.BrandC).sum()); chk('lux hidden',N_['C15'].value,((raw.Brand!=raw.BrandC)&(raw.Seg=='Luxury')).sum())
chk('fuel relabel',N_['C16'].value,(raw.Fuel_Type!=raw.FuelC).sum()); chk('km over',N_['C25'].value,(raw.Kms_Driven>1e6).sum())
lk=np.log(raw.Kms_Driven); chk('lognorm exp',N_['C26'].value,N*(1-stats.norm.cdf((math.log(1e6)-lk.mean())/lk.std(ddof=1))))
chk('year+age',N_['C32'].value,N)
for i,c in enumerate(['Mileage_kmpl','Engine_CC','Horsepower']):
    s=raw[c]; p=s[s>0]; chk(f'tech {c} flagged',N_[f'C{27+i}'].value,(~ok(s)).sum())
G=wbv['Segment_Compare']
L=raw[raw.Seg=='Luxury'].Car_Price; Nn=raw[raw.Seg=='Non-luxury'].Car_Price
chk('mean diff',G['D4'].value,L.mean()-Nn.mean()); chk('median diff %',G['E5'].value,L.median()/Nn.median()-1)
se=math.sqrt(L.var()/len(L)+Nn.var()/len(Nn)); chk('SE diff',G['B11'].value,se)
brands=sorted(raw.BrandC.unique())
for i,b in enumerate(brands): chk(f'brand median {b}',G[f'E{17+i}'].value,raw[raw.BrandC==b].Car_Price.median())
def oci(x):
    x=np.sort(np.asarray(x)); n=len(x); k=1.96*math.sqrt(n)/2; return x[max(int(n/2-k),0)], x[min(int(n/2+k),n-1)]
for i,b in enumerate(brands):
    lo,hi=oci(raw[raw.BrandC==b].Car_Price); chk(f'brand CI lo {b}',G[f'G{17+i}'].value,lo); chk(f'brand CI hi {b}',G[f'H{17+i}'].value,hi)
for i,sg in enumerate(['Luxury','Non-luxury']):
    lo,hi=oci(raw[raw.Seg==sg].Car_Price); chk(f'seg CI lo {sg}',G[f'I{4+i}'].value,lo); chk(f'seg CI hi {sg}',G[f'J{4+i}'].value,hi)
A=wbv['Age_Profile']
for a in [1,13,25]: chk(f'age {a} median',A[f'D{3+a}'].value,raw[raw.Registration_Age==a].Car_Price.median()); chk(f'age {a} lux med',A[f'F{3+a}'].value,raw[(raw.Registration_Age==a)&(raw.Seg=='Luxury')].Car_Price.median())
K=wbv['Km_Deciles']; kk=raw[raw.KMOK].copy(); kk['dec']=pd.qcut(kk.Kms_Driven,10,labels=False)
for dd in range(10):
    chk(f'decile {dd+1} n',K[f'D{4+dd}'].value,(kk.dec==dd).sum()); chk(f'decile {dd+1} median',K[f'F{4+dd}'].value,kk[kk.dec==dd].Car_Price.median())
    chk(f'decile {dd+1} lux median',K[f'H{4+dd}'].value,kk[(kk.dec==dd)&(kk.Seg=='Luxury')].Car_Price.median())
Su=wbv['Surface']; kk['ab']=((kk.Registration_Age-1)//5)
for ab,dd in [(0,0),(2,4),(4,9)]: chk(f'surface {ab},{dd}',Su.cell(4+ab,3+dd).value,kk[(kk.ab==ab)&(kk.dec==dd)].Car_Price.median())
D=wbv['Dup_Check']
key=['BrandC','Model','Year','FuelC','Transmission','Owner_Type','Color','City','Insurance_Valid','Service_History','Accidents','Tax_Paid','Number_of_Doors','Seats','Registration_Age']
cnt=raw.groupby(key).size(); chk('dup repeat rows',D['J4'].value,int((cnt-1).sum())); chk('dup pairs',D['J5'].value,int((cnt*(cnt-1)//2).sum()))
chk('price repeats',D['J6'].value,int((raw.Car_Price.value_counts()>1).sum()))
q=1.0
for b in ['Model','Year','FuelC','Transmission','Owner_Type','Color','City','Insurance_Valid','Service_History','Accidents','Tax_Paid','Number_of_Doors','Seats']: q*=float((raw[b].value_counts(normalize=True)**2).sum())
chk('prob match',D['J7'].value,q); chk('exp pairs',D['J8'].value,N*(N-1)/2*q)
first=raw.drop_duplicates(subset=key); chk('median first-of-key',D['J11'].value,first.Car_Price.median())
Mo=wbv['Mode_Bins']
pb=(np.floor(raw.Car_Price/50000)*50000); chk('price modal class',Mo['C106'].value,float(pb[pb<5e6].value_counts().index[0]+25000))
P=wbv['Plot_Sample']; 
print('sample r age',P['O5'].value,'km',P['O6'].value)
Se=wbv['Sensitivity']
u=raw[~((raw.FuelC=='Unknown')|(raw.Transmission=='Unknown')|(raw.Color=='Unknown')|(raw.City=='Unknown'))]
chk('sens unknown median',Se['C5'].value,u.Car_Price.median()); chk('sens unknown rows',Se['B5'].value,len(u))

# ---- round-2 cells
hpok=ok(raw.Horsepower); engok=ok(raw.Engine_CC)
milok=ok(raw.Mileage_kmpl)
extra=[raw.Kms_Driven.max(),0,0,0,0,0,(hpok&(raw.Horsepower<50)).sum(),(engok&(raw.Engine_CC<850)).sum(),(raw.Car_Price<20000).sum(),(hpok&(raw.Horsepower>190)).sum(),(engok&(raw.Engine_CC>2250)).sum(),(milok&(raw.Mileage_kmpl<11)).sum(),(milok&(raw.Mileage_kmpl>25.5)).sum(),raw.Horsepower[hpok].min()]
for i,e in enumerate(extra): chk(f'notes extra {47+i}',N_[f'C{47+i}'].value,float(e))
ms=raw.groupby(['Seg','Registration_Age']).Car_Price.median().unstack(0)
chk('age lux min',A['F31'].value,ms['Luxury'].min()); chk('age lux max',A['F32'].value,ms['Luxury'].max())
chk('age non min',A['G31'].value,ms['Non-luxury'].min()); chk('age non max',A['G32'].value,ms['Non-luxury'].max())
r0=[c.row for c in G['A'] if c.value=='Model'][0]; med=raw.Car_Price.median(); cov=0
for r in range(r0,r0+raw.Model.nunique()):
    b=G[f'B{r}'].value; p=raw.Car_Price[raw.Model==b]
    chk(f'model n {b}',G[f'C{r}'].value,len(p)); chk(f'model median {b}',G[f'D{r}'].value,p.median())
    lo,hi=oci(p); chk(f'model lo {b}',G[f'F{r}'].value,lo); chk(f'model hi {b}',G[f'G{r}'].value,hi); cov+=int(lo<=med<=hi)
rr=r0+raw.Model.nunique(); chk('model covering',G[f'H{rr}'].value,cov)
mm=raw.groupby('Model').Car_Price.median(); chk('model spread',G[f'D{rr}'].value,mm.max()/mm.min()-1)
bad=[c for c in checks if not c[3]]
print(f'CHECKED {len(checks)} cells; mismatches {len(bad)}')
for b in bad[:40]: print('  MISMATCH',b)
